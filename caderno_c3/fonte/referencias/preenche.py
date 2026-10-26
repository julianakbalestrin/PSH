"""Preenche a planilha PSH_Controle_Referencias_Normativas_MOP com as normas citadas na
Atividade 3.2.2 do Caderno C3 (aba Cadastro normativo) e as suas ocorrências (aba Citações no MOP).

Uso: python3 preenche.py original.xlsx saida.xlsx
"""
import copy
import datetime as dt
import math
import sys

import openpyxl
from openpyxl.styles import Alignment

from ocorrencias import levantar

ORIG, SAIDA = sys.argv[1], sys.argv[2]
HOJE = dt.datetime(2026, 10, 26)
RESP = "A designar – IAT (Atividade 3.2.2)"
CONSULTA = "Consulta preliminar em 23/10/2026, em fontes secundárias; o acesso aos portais oficiais não estava disponível."

PLANALTO = "https://www.planalto.gov.br/ccivil_03/"

# ID, esfera, tipo, número, ano, data de publicação, ementa, órgão, tema, aplicação, situação, verificação,
# URL oficial, URL consolidado, observações
NORMAS = [
    ("N-001", "Internacional", "Outro", "", "", None,
     "Acordo de Empréstimo entre o Estado do Paraná e o BIRD para o Programa de Segurança Hídrica do Paraná (PSH-PR). "
     "Parte 3.2(b): sistemas coletivos de abastecimento de água no meio rural e soluções de saneamento associadas; "
     "item (j)(vii) das definições: identificação das áreas rurais da Parte 3.2 no MOP.",
     "Estado do Paraná e Banco Internacional para Reconstrução e Desenvolvimento (BIRD)", "Governança do Programa",
     "Define o escopo da Atividade 3.2.2 e exige que o MOP identifique as áreas rurais atendidas (item 2 das Ações "
     "3.2.2.1, 3.2.2.3, 3.2.2.4 e 3.2.2.6).",
     "A verificar", "Não conferido", "", "",
     "Informar o número e a data de assinatura do Acordo. O item (j)(vii) foi citado a partir da tradução das "
     "definições usada na elaboração do MOP; conferir na versão assinada."),
    ("N-002", "Internacional", "Política / diretriz", "", "2016", None,
     "Quadro Ambiental e Social (QAS) do Banco Mundial, com as Normas Ambientais e Sociais (NAS) 1 a 10.",
     "Banco Mundial", "Ambiental e social",
     "Item 7 de todas as ações: NAS aplicáveis e medidas, com remissão aos instrumentos do Caderno AS (MGAS, PEPI, "
     "MQR, Procedimentos de Gestão de Mão de Obra e Marco de Política de Reassentamento).",
     "Vigente", "Pendente de confirmação",
     "https://www.worldbank.org/en/projects-operations/environmental-and-social-framework", "",
     "Aprovado em 2016 e aplicado aos projetos aprovados desde 1º/10/2018. A NAS 9 não se aplica ao Programa. "
     "Conferir a versão em português adotada no Caderno AS."),
    ("N-003", "Internacional", "Política / diretriz", "", "", None,
     "Regulamento de Aquisições para Mutuários de Financiamento de Projetos de Investimento (Procurement "
     "Regulations for IPF Borrowers).",
     "Banco Mundial", "Aquisições",
     "Contratações financiadas pelo BIRD nas Ações 3.2.2.1, 3.2.2.2, 3.2.2.3, 3.2.2.5, 3.2.2.6 e 3.2.2.7, conforme o "
     "Caderno AQ e o Plano de Aquisições; compatibilidade do regime de contratação integrada (Ação 3.2.2.3).",
     "A verificar", "Pendente de confirmação",
     "https://www.worldbank.org/en/projects-operations/products-and-services/brief/procurement-new-framework", "",
     "Vale a edição indicada no Acordo de Empréstimo. " + CONSULTA + " Foram encontradas a 6ª edição (fev. 2025, "
     "em vigor desde 1º/3/2025) e a 7ª edição (set. 2025)."),
    ("N-004", "Federal", "Lei", "9.433", "1997", dt.datetime(1997, 1, 9),
     "Lei nº 9.433, de 8 de janeiro de 1997. Institui a Política Nacional de Recursos Hídricos, cria o Sistema "
     "Nacional de Gerenciamento de Recursos Hídricos, regulamenta o inciso XIX do art. 21 da Constituição Federal, e "
     "altera o art. 1º da Lei nº 8.001, de 13 de março de 1990, que modificou a Lei nº 7.990, de 28 de dezembro de 1989.",
     "Presidência da República", "Recursos hídricos – outorga",
     "Outorga de direito de uso da água subterrânea dos poços (Ação 3.2.2.3, item 10).",
     "Alterada / consolidada", "Pendente de confirmação", PLANALTO + "leis/l9433.htm", PLANALTO + "leis/l9433.htm",
     "Conferir o texto compilado no Planalto."),
    ("N-005", "Federal", "Lei", "11.107", "2005", dt.datetime(2005, 4, 7),
     "Lei nº 11.107, de 6 de abril de 2005. Dispõe sobre normas gerais de contratação de consórcios públicos e dá "
     "outras providências.",
     "Presidência da República", "Gestão associada – consórcios públicos",
     "Alternativa de modelo de gestão por consórcio público, com alteração do contrato de consórcio, contrato de "
     "rateio e cessão de uso dos ativos (Ação 3.2.2.1, Etapa 1).",
     "Alterada / consolidada", "Pendente de confirmação",
     PLANALTO + "_ato2004-2006/2005/lei/l11107.htm", PLANALTO + "_ato2004-2006/2005/lei/l11107.htm",
     "Alterada, entre outras, pela Lei nº 14.026/2020, que vedou a prestação por contrato de programa dos serviços de "
     "que trata o art. 175 da Constituição. Conferir o texto compilado."),
    ("N-006", "Federal", "Lei", "11.445", "2007", dt.datetime(2007, 1, 8),
     "Lei nº 11.445, de 5 de janeiro de 2007. Estabelece as diretrizes nacionais para o saneamento básico [...].",
     "Presidência da República", "Saneamento básico – marco legal",
     "Titularidade, gestão associada e prestação dos serviços de saneamento no meio rural; base do arranjo "
     "consorciado e das regras de tarifa (Ação 3.2.2.1).",
     "Alterada / consolidada", "Pendente de confirmação",
     PLANALTO + "_ato2007-2010/2007/lei/l11445.htm", PLANALTO + "_ato2007-2010/2007/lei/l11445.htm",
     "Redação atualizada pela Lei nº 14.026/2020. Atenção: a nota técnica do Componente 3 traz trocados, na lista de "
     "referências, os endereços das Leis nº 11.445/2007 e nº 14.026/2020."),
    ("N-007", "Federal", "Lei", "13.019", "2014", dt.datetime(2014, 8, 1),
     "Lei nº 13.019, de 31 de julho de 2014. Estabelece o regime jurídico das parcerias entre a administração pública "
     "e as organizações da sociedade civil [...] (Marco Regulatório das Organizações da Sociedade Civil – MROSC).",
     "Presidência da República", "Parcerias com organizações da sociedade civil",
     "Alternativas de gestão por associação de moradores e por central de associações; Acordo de Cooperação para o "
     "uso da infraestrutura pública (Ação 3.2.2.1).",
     "Alterada / consolidada", "Pendente de confirmação",
     PLANALTO + "_ato2011-2014/2014/lei/l13019.htm", PLANALTO + "_ato2011-2014/2014/lei/l13019.htm",
     "Alterada pela Lei nº 13.204/2015, entre outras. Conferir o texto compilado e a regulamentação estadual."),
    ("N-008", "Federal", "Lei", "13.709", "2018", dt.datetime(2018, 8, 15),
     "Lei nº 13.709, de 14 de agosto de 2018. Lei Geral de Proteção de Dados Pessoais (LGPD).",
     "Presidência da República", "Proteção de dados pessoais",
     "Tratamento dos dados pessoais do cadastro domiciliar do trabalho técnico-social (Ação 3.2.2.2).",
     "Alterada / consolidada", "Pendente de confirmação",
     PLANALTO + "_ato2015-2018/2018/lei/l13709.htm", PLANALTO + "_ato2015-2018/2018/lei/l13709.htm",
     "Alterada pela Lei nº 13.853/2019, entre outras. Conferir o texto compilado."),
    ("N-009", "Federal", "Lei", "14.026", "2020", dt.datetime(2020, 7, 16),
     "Lei nº 14.026, de 15 de julho de 2020. Atualiza o marco legal do saneamento básico e altera a Lei nº 9.984, de "
     "17 de julho de 2000, [...].",
     "Presidência da República", "Saneamento básico – marco legal",
     "Marco legal atualizado do saneamento: limita o contrato de programa e trata da prestação regionalizada, o que "
     "afeta o arranjo consorciado (Ação 3.2.2.1).",
     "Vigente", "Pendente de confirmação",
     PLANALTO + "_ato2019-2022/2020/lei/l14026.htm", "",
     "Lei de alteração: as regras vigentes estão nos textos compilados das Leis nº 11.445/2007 e nº 11.107/2005."),
    ("N-010", "Federal", "Lei", "14.133", "2021", dt.datetime(2021, 4, 1),
     "Lei nº 14.133, de 1º de abril de 2021. Lei de Licitações e Contratos Administrativos.",
     "Presidência da República", "Aquisições",
     "Definição do regime de contratação integrada ou semi-integrada (art. 6º, XXXIII) das obras de água (Ação "
     "3.2.2.3, Etapa 3).",
     "A verificar", "Pendente de confirmação",
     PLANALTO + "_ato2019-2022/2021/lei/l14133.htm", PLANALTO + "_ato2019-2022/2021/lei/l14133.htm",
     "Citada em comentário de revisão, não no texto. Nas contratações financiadas pelo BIRD prevalece o Regulamento de "
     "Aquisições do Banco (N-003)."),
    ("N-011", "Federal", "Decreto", "11.599", "2023", dt.datetime(2023, 7, 13),
     "Decreto nº 11.599, de 12 de julho de 2023. Dispõe sobre a prestação regionalizada dos serviços públicos de "
     "saneamento básico, o apoio técnico e financeiro de que trata o art. 13 da Lei nº 14.026, de 15 de julho de "
     "2020, a alocação de recursos públicos federais e os financiamentos com recursos da União [...].",
     "Presidência da República", "Saneamento básico – prestação regionalizada",
     "Citado pela nota técnica do Componente 3 como respaldo aos arranjos associativos (Ação 3.2.2.1, Etapa 1, e "
     "quadro de alternativas de gestão).",
     "A verificar", "Pendente de confirmação",
     PLANALTO + "_ato2023-2026/2023/decreto/d11599.htm", PLANALTO + "_ato2023-2026/2023/decreto/d11599.htm",
     "Conferir a aderência: o decreto trata da prestação regionalizada e do apoio técnico e financeiro da União, não "
     "de associações comunitárias. Conferir alterações posteriores."),
    ("N-012", "Federal", "Portaria", "2.777", "2025", dt.datetime(2025, 8, 5),
     "Portaria nº 2.777, de 30 de julho de 2025, da Funasa. Segundo o portal da Funasa, trata da seleção de propostas "
     "de sistemas de abastecimento de água, de esgotamento sanitário e de melhorias sanitárias domiciliares em áreas "
     "rurais (ementa a transcrever do Diário Oficial da União).",
     "Ministério da Saúde – Fundação Nacional de Saúde (Funasa)", "Saneamento rural – investimentos",
     "Citada pela nota técnica do Componente 3 como diretriz que valida o modelo de centrais e do SISAR (Ação "
     "3.2.2.1, Etapa 1, e quadro de alternativas de gestão).",
     "A verificar", "Pendente de confirmação",
     "https://bvs.saude.gov.br/bvs/saudelegis/funasa/2025/prt2777_05_08_2025.html", "",
     "Endereço indicado na nota técnica; a data de publicação foi inferida desse endereço. " + CONSULTA + " O "
     "portal da Funasa a apresenta como seleção de propostas, com alterações posteriores, e não como diretriz de "
     "modelo de gestão: se confirmado, ajustar o texto do MOP."),
    ("N-013", "Federal", "Portaria", "888", "2021", dt.datetime(2021, 5, 7),
     "Portaria GM/MS nº 888, de 4 de maio de 2021. Altera o Anexo XX da Portaria de Consolidação GM/MS nº 5, de 28 de "
     "setembro de 2017, para dispor sobre os procedimentos de controle e de vigilância da qualidade da água para "
     "consumo humano e seu padrão de potabilidade.",
     "Ministério da Saúde", "Qualidade da água – potabilidade",
     "Padrão de potabilidade nos testes de entrega, nos critérios de aceite e no acompanhamento da água distribuída "
     "(Ações 3.2.2.3 e 3.2.2.4).",
     "A verificar", "Pendente de confirmação",
     "https://bvsms.saude.gov.br/bvs/saudelegis/gm/2021/prt0888_07_05_2021.html",
     "https://bvsms.saude.gov.br/bvs/saudelegis/gm/2017/prc0005_03_10_2017.html",
     "O texto vigente está no Anexo XX da Portaria de Consolidação nº 5/2017 (N-014). Fonte secundária indica "
     "alteração pela Portaria GM/MS nº 2.472/2021; conferir."),
    ("N-014", "Federal", "Portaria", "5", "2017", dt.datetime(2017, 10, 3),
     "Portaria de Consolidação GM/MS nº 5, de 28 de setembro de 2017. Consolidação das normas sobre as ações e os "
     "serviços de saúde do SUS; o Anexo XX trata do controle e da vigilância da qualidade da água para consumo humano.",
     "Ministério da Saúde", "Qualidade da água – potabilidade",
     "Texto consolidado do padrão de potabilidade (Anexo XX), citado na Ação 3.2.2.3, item 10.",
     "Alterada / consolidada", "Pendente de confirmação",
     "https://bvsms.saude.gov.br/bvs/saudelegis/gm/2017/prc0005_03_10_2017.html",
     "https://bvsms.saude.gov.br/bvs/saudelegis/gm/2017/prc0005_03_10_2017.html",
     "Citar no MOP o Anexo XX, com a redação dada pela Portaria GM/MS nº 888/2021."),
    ("N-015", "Federal", "Portaria", "3.174", "2019", dt.datetime(2019, 12, 4),
     "Portaria nº 3.174, de 2 de dezembro de 2019, do Ministério da Saúde. Dispõe sobre o Programa Nacional de "
     "Saneamento Rural (PNSR).",
     "Ministério da Saúde", "Saneamento rural – política nacional",
     "Referência para os eixos de gestão, educação e participação social e tecnologia (Ações 3.2.2.1 e 3.2.2.2).",
     "A verificar", "Pendente de confirmação", "", "",
     "O MOP cita o documento do PNSR (BRASIL, 2019, Funasa); o número e a data da portaria vêm de fontes "
     "secundárias. Localizar no Saúde Legis."),
    ("N-016", "Outra", "Norma técnica", "17076", "2024", None,
     "ABNT NBR 17076:2024 – Projeto de sistema de tratamento de esgoto de menor porte – Requisitos.",
     "Associação Brasileira de Normas Técnicas (ABNT)", "Esgotamento sanitário – soluções individuais",
     "Projeto, instalação e aceite das soluções individuais de esgotamento sanitário (Ações 3.2.2.5 e 3.2.2.6).",
     "Vigente", "Pendente de confirmação", "https://www.abntcatalogo.com.br/", "",
     "Publicada em abr. 2024. Aplica-se a vazões de até 12.000 L/dia e carga orgânica de até 3,80 kg DBO/dia. "
     "Catálogo de terceiros registra errata ou nova edição em 05/2025; conferir a edição vigente no Catálogo ABNT."),
    ("N-017", "Outra", "Norma técnica", "7229", "1993", None,
     "ABNT NBR 7229:1993 – Projeto, construção e operação de sistemas de tanques sépticos.",
     "Associação Brasileira de Normas Técnicas (ABNT)", "Esgotamento sanitário – soluções individuais",
     "Citada como norma substituída pela ABNT NBR 17076:2024 (Ação 3.2.2.5, item 1).",
     "Revogada", "Pendente de confirmação", "https://www.abntcatalogo.com.br/", "",
     "Cancelada e substituída pela NBR 17076:2024 (confirmar no Catálogo ABNT); não usar em projetos."),
    ("N-018", "Outra", "Norma técnica", "13969", "1997", None,
     "ABNT NBR 13969:1997 – Tanques sépticos – Unidades de tratamento complementar e disposição final dos efluentes "
     "líquidos – Projeto, construção e operação.",
     "Associação Brasileira de Normas Técnicas (ABNT)", "Esgotamento sanitário – soluções individuais",
     "Citada como norma substituída pela ABNT NBR 17076:2024 (Ação 3.2.2.5, item 1).",
     "Revogada", "Pendente de confirmação", "https://www.abntcatalogo.com.br/", "",
     "Cancelada e substituída pela NBR 17076:2024 (confirmar no Catálogo ABNT); não usar em projetos."),
    ("N-019", "Outra", "Norma técnica", "12212", "2017", None,
     "ABNT NBR 12212:2017 – Poço tubular – Projeto de poço tubular para captação de água subterrânea – Procedimento.",
     "Associação Brasileira de Normas Técnicas (ABNT)", "Abastecimento de água – poços",
     "Projeto dos poços tubulares profundos (Ação 3.2.2.3, item 10).",
     "Vigente", "Pendente de confirmação", "https://www.abntcatalogo.com.br/", "",
     "Publicada em set. 2017; catálogo de terceiros registra confirmação da edição em 01/2022. Conferir no Catálogo ABNT."),
    ("N-020", "Outra", "Norma técnica", "12244", "2006", None,
     "ABNT NBR 12244:2006 – Poço tubular – Construção de poço tubular para captação de água subterrânea.",
     "Associação Brasileira de Normas Técnicas (ABNT)", "Abastecimento de água – poços",
     "Construção dos poços tubulares profundos (Ação 3.2.2.3, item 10).",
     "Vigente", "Pendente de confirmação", "https://www.abntcatalogo.com.br/", "",
     "Publicada em mar. 2006. Conferir no Catálogo ABNT."),
    ("N-021", "Estadual - PR", "Lei", "12.726", "1999", None,
     "Lei nº 12.726, de 26 de novembro de 1999. Institui a Política Estadual de Recursos Hídricos e adota outras "
     "providências.",
     "Assembleia Legislativa e Governo do Estado do Paraná", "Recursos hídricos – outorga",
     "Outorga estadual de direito de uso da água dos poços, emitida pelo IAT (Ação 3.2.2.3, item 10).",
     "A verificar", "Pendente de confirmação", "https://www.legislacao.pr.gov.br/", "",
     "Conferir o texto compilado e a regulamentação da outorga no Sistema Estadual de Legislação."),
    ("N-022", "Estadual - PR", "Outro", "", "", None,
     "Regulamento de Licitações e Contratos da Companhia de Saneamento do Paraná (Sanepar).",
     "Sanepar", "Aquisições – contrapartida",
     "Aquisições feitas pela Sanepar com recursos de contrapartida no Programa Sanepar Rural (Ação 3.2.2.4).",
     "A verificar", "Não conferido", "", "",
     "Informar a versão vigente (base: Lei nº 13.303/2016). Confirmar com a UGP e o Banco Mundial as regras aplicáveis "
     "às despesas de contrapartida."),
]

FINALIDADE = {
    "N-001": "Vincular a atividade ao Acordo e à exigência de identificar os locais",
    "N-002": "Indicar as NAS aplicáveis e as medidas da ação",
    "N-003": "Indicar as regras das contratações financiadas pelo BIRD",
    "N-004": "Fundamentar a outorga de uso da água",
    "N-005": "Fundamentar a alternativa de gestão por consórcio público",
    "N-006": "Situar o arranjo de gestão no marco legal do saneamento",
    "N-007": "Fundamentar a gestão por associações e o Acordo de Cooperação",
    "N-008": "Orientar o tratamento dos dados pessoais do cadastro",
    "N-009": "Situar os arranjos no marco legal atualizado do saneamento",
    "N-010": "Definir o regime de contratação das obras",
    "N-011": "Respaldo aos arranjos associativos, segundo a nota técnica",
    "N-012": "Respaldo aos arranjos associativos, segundo a nota técnica",
    "N-013": "Definir o padrão de potabilidade da água entregue",
    "N-014": "Indicar o texto consolidado do padrão de potabilidade",
    "N-015": "Referência conceitual do saneamento rural",
    "N-016": "Definir os requisitos de projeto e instalação das soluções individuais",
    "N-017": "Registrar a norma substituída pela NBR 17076:2024",
    "N-018": "Registrar a norma substituída pela NBR 17076:2024",
    "N-019": "Definir os requisitos de projeto dos poços",
    "N-020": "Definir os requisitos de construção dos poços",
    "N-021": "Fundamentar a outorga estadual de uso da água",
    "N-022": "Indicar as regras de aquisição da Sanepar",
}
VINCULO = {
    "N-001": "Fundamento legal", "N-002": "Requisito ambiental e social", "N-003": "Aquisições / contratos",
    "N-004": "Fundamento legal", "N-005": "Fundamento legal", "N-006": "Fundamento legal",
    "N-007": "Fundamento legal", "N-008": "Regra de execução", "N-009": "Fundamento legal",
    "N-010": "Aquisições / contratos", "N-011": "Fundamento legal", "N-012": "Fundamento legal",
    "N-013": "Regra de execução", "N-014": "Regra de execução", "N-015": "Referência técnica",
    "N-016": "Referência técnica", "N-017": "Referência técnica", "N-018": "Referência técnica",
    "N-019": "Referência técnica", "N-020": "Referência técnica", "N-021": "Fundamento legal",
    "N-022": "Aquisições / contratos",
}


def altura(ws, valores, larguras, minimo=23.0):
    """Altura da linha estimada pelo texto mais longo (Carlito 11, quebra automática)."""
    linhas = 1
    for col, v in valores.items():
        if v is None:
            continue
        cpl = max(4, larguras.get(col, 10) * 1.15)
        n = sum(max(1, math.ceil(len(par) / cpl)) for par in str(v).split("\n"))
        linhas = max(linhas, n)
    return max(minimo, 15.0 * linhas + 6)


def escreve(ws, linha0, linhas, colunas, modelo_linha=6):
    larg = {k: v.width for k, v in ws.column_dimensions.items()}
    for i, dados in enumerate(linhas):
        r = linha0 + i
        valores = {}
        for j, v in enumerate(dados):
            col = colunas[j]
            c = ws[f"{col}{r}"]
            mod = ws[f"{col}{modelo_linha}"]
            c.font = copy.copy(mod.font)
            c.border = copy.copy(mod.border)
            c.fill = copy.copy(mod.fill)
            c.number_format = mod.number_format
            c.protection = copy.copy(mod.protection)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.value = v if v != "" else None
            valores[col] = v
        ws.row_dimensions[r].height = altura(ws, valores, larg)


def main():
    wb = openpyxl.load_workbook(ORIG)

    # ------------------------------------------------------------ cadastro normativo
    ws = wb["Cadastro normativo"]
    cols = list("ABCDEFGHIJKLMNOPQ")
    linhas = []
    for (nid, esf, tipo, num, ano, data, ementa, orgao, tema, aplic, sit, verif, url, urlc, obs) in NORMAS:
        linhas.append([nid, esf, tipo, num, ano, data, ementa, orgao, tema, aplic, sit, verif, url, urlc,
                       RESP, None, obs])
    escreve(ws, 5, linhas, cols)
    t = ws.tables["NormasPSH"]
    t.ref = f"A4:Q{4 + len(linhas)}"

    # ------------------------------------------------------------ citações no MOP
    ws2 = wb["Citações no MOP"]
    oc = levantar()
    ids = {n[0] for n in NORMAS}
    linhas2 = []
    for k, o in enumerate(oc, 1):
        assert o["norma"] in ids, o["norma"]
        if o["acao"] == "Atividade 3.2.2":
            secao, acao = "Atividade 3.2.2 – apresentação", "Atividade 3.2.2"
        else:
            secao = o["item"]
            if o["etapa"]:
                secao += " › " + o["etapa"]
            acao = "Ação " + o["acao"]
        if o["rotulo"]:
            secao += " › " + o["rotulo"]
        obs = {"texto": "", "quadro": "Citação em quadro.", "figura": "Citação em figura.",
               "comentário": "Citação em comentário de revisão do Word, não no corpo do texto."}[o["grupo"]]
        if o["n"] > 1:
            obs = (obs + f" Mais {o['n'] - 1} menção(ões) no mesmo item.").strip()
        precisa, status = "A avaliar", "Pendente"
        if o["norma"] in ("N-011", "N-012"):
            status = "Em revisão"
            obs = (obs + " Conferir se a norma dá o respaldo citado pela nota técnica (ver Cadastro normativo).").strip()
        if o["norma"] in ("N-017", "N-018"):
            precisa = "Não"
        if o["norma"] == "N-001" and o["grupo"] == "comentário":
            precisa = "Sim"
            obs = (obs + " Incluir no MOP a lista ou o mapa dos locais de intervenção.").strip()
        linhas2.append([f"C-{k:03d}", o["norma"], "Componente 3", secao, acao, o["trecho"], FINALIDADE[o["norma"]],
                        VINCULO[o["norma"]], precisa, "IAT – equipe da Atividade 3.2.2", status, HOJE, obs])
    escreve(ws2, 5, linhas2, list("ABCDEFGHIJKLM"))
    t2 = ws2.tables["CitacoesPSH"]
    t2.ref = f"A4:M{4 + len(linhas2)}"

    # ------------------------------------------------------------ painel: as linhas-modelo foram substituídas
    wp = wb["Painel"]
    wp["B5"] = "=COUNTA('Cadastro normativo'!A5:A300)"
    wp["B6"] = "=COUNTA('Citações no MOP'!A5:A500)"
    wp["B7"] = "=COUNTIF('Cadastro normativo'!L5:L300,\"Conferido em fonte oficial\")"
    wp["B8"] = ("=COUNTIF('Citações no MOP'!K5:K500,\"Pendente\")+COUNTIF('Citações no MOP'!K5:K500,"
                "\"Ajustar texto\")")
    wp["A10"] = ("As linhas-modelo 5 foram substituídas pelo cadastro da Atividade 3.2.2 (Caderno C3) em "
                 "26/10/2026; os indicadores contam todas as linhas preenchidas.")

    wb.save(SAIDA)
    print(f"normas: {len(linhas)} | citações: {len(linhas2)} -> {SAIDA}")


main()
